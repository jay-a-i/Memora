import logging
import os
import shutil
import docx
from typing import Dict

import pymupdf
import pymupdf4llm

from backend.app.core.config import settings

logger = logging.getLogger(__name__)


def doc_to_md(
    file_path: str,
    output_md: str = "temp_output.md",
    image_dir: str = "temp_image_dir",
    batch_size: int = 50,
    extract_images: bool | None = None,
) -> Dict[str, str] | None:
    """Converts a PDF file into Markdown format in memory-efficient page batches.

    Args:
        file_path (str): Path to the source PDF file.
        output_md (str, optional): Target file path for the output Markdown text.
        batch_size (int, optional): Number of PDF pages to process per batch to manage
            RAM usage. Defaults to 50.
        extract_images (bool, optional): Write embedded images to disk. Defaults to
            EXTRACT_IMAGES. Images are not ingested, so this is off by default;
            enabling it only fills disk with files that are deleted after upload.

    Returns:
        Dict[str, str]: A dictionary containing file paths for 'output_md' and 'image_dir' if successful; otherwise None.
    """
    if extract_images is None:
        extract_images = settings.EXTRACT_IMAGES

    if os.path.exists(output_md):
        os.remove(output_md)
    if os.path.exists(image_dir):
        shutil.rmtree(image_dir)

    try:
        with pymupdf.open(file_path) as doc:
            if doc.is_encrypted and not doc.authenticate(""):
                logger.error("PDF is encrypted and cannot be opened: %s", file_path)
                return None

            total_pages = len(doc)
            if total_pages == 0:
                logger.error("PDF has no pages: %s", file_path)
                return None

            with open(output_md, "w", encoding="utf-8") as f:
                for i in range(0, total_pages, batch_size):
                    end_page = min(i + batch_size, total_pages)
                    pages_to_extract = list(range(i, end_page))

                    md_batch = pymupdf4llm.to_markdown(
                        doc,
                        pages=pages_to_extract,
                        write_images=extract_images,
                        image_path=image_dir if extract_images else None,
                        force_ocr=False,
                    )
                    f.write(md_batch + "\n\n")
                    logger.info("Converted pages %s-%s", i, end_page - 1)

        return {
            "output_md": output_md,
            "image_dir": image_dir
        }

    except Exception:
        logger.exception("PDF to Markdown conversion failed for %s", file_path)
        return None


def _iter_body_blocks(document):
    """
    Yields the document body in true reading order.

    `document.paragraphs` and `document.tables` are two independent flat lists
    with no positional information, so walking one and then the other moved
    every table to the end of the document. A table under "# Intro" was emitted
    after "# Appendix", and since chunking splits on headers, that table's chunk
    was filed under the wrong section -- a metadata error invisible in the
    stored text.

    Iterating the XML body children preserves the interleaving, because a
    paragraph and a table are siblings in the same element sequence.
    """
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    body = document.element.body
    for child in body.iterchildren():
        # Skip the section properties element; it carries no content.
        tag = child.tag.split("}")[-1]
        if tag == "p":
            yield _paragraph_to_markdown(Paragraph(child, document))
        elif tag == "tbl":
            yield Table(child, document)


def _paragraph_to_markdown(para) -> str:
    """Maps one paragraph onto a Markdown line, honouring headings and lists."""
    text = para.text.strip()
    if not text:
        return ""

    style = (para.style.name or "").lower() if para.style else ""
    if style.startswith("heading"):
        level = style.replace("heading", "").strip()
        if level.isdigit() and 1 <= int(level) <= 6:
            return f"{'#' * int(level)} {text}"
        return text

    # Word's list styles carried no marker before, so bulleted and numbered
    # items were flattened into bare prose and lost their structure cue.
    if style.startswith("list"):
        if "number" in style:
            return f"1. {text}"
        return f"- {text}"

    return text


def _escape_cell(text: str) -> str:
    """
    Makes a value safe to place inside a Markdown table cell.

    A raw `|` in the source produced a row with more cells than the header, and
    an embedded newline split one row across two. Both corrupt the table and
    cause the splitter to cut through its middle.
    """
    return text.replace("\\", "\\\\").replace("|", "\\|").replace("\n", " ").strip()


def _table_to_markdown(table) -> list[str]:
    """Renders one table as a Markdown table block."""
    rows = [
        [_escape_cell(cell.text) for cell in row.cells] for row in table.rows
    ]
    rows = [r for r in rows if any(cell for cell in r)]
    if not rows:
        return []

    width = max(len(r) for r in rows)
    padded = [r + [""] * (width - len(r)) for r in rows]

    lines = [""]
    lines.append("| " + " | ".join(padded[0]) + " |")
    lines.append("| " + " | ".join("---" for _ in range(width)) + " |")
    for row in padded[1:]:
        lines.append("| " + " | ".join(row) + " |")
    lines.append("")
    return lines


def docx_to_md(
    file_path: str,
    output_md: str = "temp_output.md",
) -> Dict[str, str] | None:
    """Converts a .docx file to Markdown, preserving headings, lists and tables.

    A .docx is a ZIP archive, so reading it as text would embed compressed
    binary noise. This walks the document body in order and maps Word's
    paragraph styles onto Markdown headings and list markers.

    Args:
        file_path (str): Path to the source .docx file.
        output_md (str, optional): Target path for the Markdown output.

    Returns:
        Dict[str, str] | None: {'output_md': path} on success, else None.
    """
    try:
        document = docx.Document(file_path)

        lines: list[str] = []
        for block in _iter_body_blocks(document):
            if isinstance(block, str):
                lines.append(block)
            else:
                lines.extend(_table_to_markdown(block))

        markdown = "\n\n".join(
            "\n".join(lines).split("\n\n\n")
        ).strip()

        if not markdown:
            logger.error("No extractable text in %s", file_path)
            return None

        with open(output_md, "w", encoding="utf-8") as f:
            f.write(markdown + "\n")

        return {"output_md": output_md}

    except Exception:
        logger.exception("DOCX to Markdown conversion failed for %s", file_path)
        return None
